"""Compara el exportador BCP legacy y el nuevo sobre el mismo PDF sintético."""

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from shutil import rmtree
from unittest import TestCase, skipIf
from uuid import uuid4

from openpyxl import load_workbook

from statement_worker.domain.models import ExtractionStatus
from statement_worker.exporters import build_bcp_workbook_plan
from statement_worker.exporters.xlsx_writer import export_workbook_plan
from statement_worker.extractors.bcp.document_processor import process_bcp_pdf
from tests.characterization.legacy_harness import run_legacy_bcp_export, unavailable_reason
from tests.support.synthetic_bcp_pdf import SYNTHETIC_HOLDER, create_synthetic_bcp_pdf

_UNAVAILABLE = unavailable_reason()

# Campos de `Encabezado` que el legacy y el nuevo leen igual (ADR-0015).
_SHARED_HEADER_FIELDS = (
    ("codigo_cuenta", "Código de cuenta"),
    ("moneda", "Moneda"),
    ("periodo", "Periodo"),
    ("pagina_inicial", "Página inicial"),
    ("total_paginas_declarado", "Total de páginas declarado"),
)


def _amount(value: object) -> Decimal | None:
    if value is None:
        return None
    return Decimal(str(value))


def _legacy_date(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    return datetime.strptime(str(value), "%d/%m/%Y").date()


def _new_date(value: object) -> object:
    if isinstance(value, datetime):
        return value.date()
    return value


def _rows(path: Path, sheet: str) -> tuple[tuple[object, ...], ...]:
    workbook = load_workbook(filename=path, data_only=True)
    try:
        return tuple(tuple(row) for row in workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()


def _columns(records: tuple[tuple[object, ...], ...]) -> dict[str, int]:
    return {str(name): index for index, name in enumerate(records[0])}


@skipIf(_UNAVAILABLE is not None, str(_UNAVAILABLE))
class LegacyBcpParityTests(TestCase):
    """El nuevo exportador debe igualar o mejorar el resultado del legacy."""

    directory: Path
    legacy_xlsx: Path
    new_xlsx: Path

    @classmethod
    def setUpClass(cls) -> None:
        cls.directory = Path.cwd() / f".test_legacy_parity_{uuid4().hex}"
        cls.directory.mkdir()
        cls.addClassCleanup(rmtree, cls.directory, True)

        pdf_path = cls.directory / "synthetic-bcp.pdf"
        create_synthetic_bcp_pdf(pdf_path)

        cls.legacy_xlsx = cls.directory / "legacy.xlsx"
        run_legacy_bcp_export(pdf_path, cls.legacy_xlsx)

        cls.result = process_bcp_pdf(pdf_path, temporary_parent=cls.directory)
        cls.plan = build_bcp_workbook_plan(cls.result)
        cls.new_xlsx = cls.directory / "nuevo.xlsx"
        export_workbook_plan(cls.plan, cls.new_xlsx)

    def test_both_exporters_produce_a_readable_statement(self) -> None:
        self.assertTrue(self.legacy_xlsx.is_file())
        self.assertTrue(self.new_xlsx.is_file())
        self.assertEqual(self.result.validation.status, ExtractionStatus.SUCCEEDED)

    def test_row_count_types_and_pages_match(self) -> None:
        legacy = _rows(self.legacy_xlsx, "Movimientos")
        new = _rows(self.new_xlsx, "Movimientos")
        legacy_columns = _columns(legacy)
        new_columns = _columns(new)

        self.assertEqual(len(legacy) - 1, len(new) - 1)
        for index, (legacy_row, new_row) in enumerate(zip(legacy[1:], new[1:], strict=True)):
            with self.subTest(row=index):
                # Los mismos rótulos que el legacy, en español (ADR-0015).
                self.assertEqual(
                    legacy_row[legacy_columns["tipo_fila"]],
                    new_row[new_columns["Tipo de fila"]],
                )
                self.assertEqual(
                    legacy_row[legacy_columns["pagina_pdf"]],
                    new_row[new_columns["Página"]],
                )

    def test_descriptions_dates_and_amounts_match(self) -> None:
        legacy = _rows(self.legacy_xlsx, "Movimientos")
        new = _rows(self.new_xlsx, "Movimientos")
        legacy_columns = _columns(legacy)
        new_columns = _columns(new)
        pairs = (
            ("descripcion", "Descripción"),
            ("cargo", "Cargo"),
            ("abono", "Abono"),
            ("saldo", "Saldo"),
        )

        for index, (legacy_row, new_row) in enumerate(zip(legacy[1:], new[1:], strict=True)):
            for legacy_key, new_key in pairs:
                with self.subTest(row=index, column=legacy_key):
                    legacy_value = legacy_row[legacy_columns[legacy_key]]
                    new_value = new_row[new_columns[new_key]]
                    if legacy_key == "descripcion":
                        self.assertEqual(legacy_value, new_value)
                    else:
                        self.assertEqual(_amount(legacy_value), _amount(new_value))
            for legacy_key, new_key in (
                ("fecha_proc", "Fecha proceso"),
                ("fecha_valor", "Fecha valor"),
            ):
                with self.subTest(row=index, column=legacy_key):
                    self.assertEqual(
                        _legacy_date(legacy_row[legacy_columns[legacy_key]]),
                        _new_date(new_row[new_columns[new_key]]),
                    )

    def test_summary_totals_match(self) -> None:
        legacy = _rows(self.legacy_xlsx, "Resumen")
        new = _rows(self.new_xlsx, "Resumen")
        legacy_columns = _columns(legacy)
        new_columns = _columns(new)
        legacy_row, new_row = legacy[1], new[1]

        for legacy_key, new_key in (
            ("total_paginas", "Páginas"),
            ("total_filas_extraidas", "Filas"),
            ("total_movimientos", "Movimientos"),
        ):
            with self.subTest(column=legacy_key):
                self.assertEqual(
                    legacy_row[legacy_columns[legacy_key]],
                    new_row[new_columns[new_key]],
                )
        for legacy_key, new_key in (
            ("total_cargos", "Total cargos"),
            ("total_abonos", "Total abonos"),
            ("ultimo_saldo_detectado", "Saldo final"),
        ):
            with self.subTest(column=legacy_key):
                self.assertEqual(
                    _amount(legacy_row[legacy_columns[legacy_key]]),
                    _amount(new_row[new_columns[new_key]]),
                )

    def test_page_control_matches(self) -> None:
        legacy = _rows(self.legacy_xlsx, "Control_Paginas")
        new = _rows(self.new_xlsx, "Control_Paginas")
        legacy_columns = _columns(legacy)
        new_columns = _columns(new)

        movements = _rows(self.new_xlsx, "Movimientos")
        page_index = _columns(movements)["Página"]
        summary = _rows(self.new_xlsx, "Resumen")

        self.assertEqual(len(legacy) - 1, len(new) - 1)
        for index, (legacy_row, new_row) in enumerate(zip(legacy[1:], new[1:], strict=True)):
            with self.subTest(page=index):
                page = new_row[new_columns["Página"]]
                self.assertEqual(legacy_row[legacy_columns["pagina_pdf"]], page)
                # El legacy cuenta aparte la segunda línea de un concepto partido; el
                # nuevo cuenta las filas que publica `Movimientos` en esa página.
                self.assertEqual(
                    new_row[new_columns["Filas extraídas"]],
                    sum(1 for row in movements[1:] if row[page_index] == page),
                )
                self.assertLessEqual(
                    new_row[new_columns["Filas extraídas"]],
                    legacy_row[legacy_columns["filas_extraidas"]],
                )
                self.assertEqual(legacy_row[legacy_columns["texto_detectado"]], "SI")
                self.assertIs(new_row[new_columns["Cabecera detectada"]], True)
        # Así la hoja suma lo mismo que el `Resumen`.
        self.assertEqual(
            sum(row[new_columns["Filas extraídas"]] for row in new[1:]),
            summary[1][_columns(summary)["Filas"]],
        )

    def test_the_new_export_improves_the_documented_weak_points(self) -> None:
        legacy_workbook = load_workbook(filename=self.legacy_xlsx, data_only=True)
        try:
            legacy_sheets = tuple(legacy_workbook.sheetnames)
            legacy_header_fields = {
                str(row[0]) for row in legacy_workbook["Encabezado"].iter_rows(values_only=True)
            }
            legacy_amount_types = {
                type(row[5]).__name__
                for row in legacy_workbook["Movimientos"].iter_rows(min_row=2, values_only=True)
                if row[5] is not None
            }
        finally:
            legacy_workbook.close()

        # Las mismas hojas que el legacy: `Encabezado` vuelve a pedido del producto
        # y las invariantes las publica la API, no una hoja del libro (ADR-0015).
        self.assertEqual(tuple(sheet.name for sheet in self.plan.sheets), legacy_sheets)
        self.assertIn("titular", legacy_header_fields)

        # El legacy escribe importes como números de Python; el nuevo conserva `Decimal`.
        self.assertNotIn("Decimal", legacy_amount_types)
        movements = next(sheet for sheet in self.plan.sheets if sheet.name == "Movimientos")
        amounts = [value for row in movements.rows for value in row[5:8] if value is not None]
        self.assertTrue(amounts)
        self.assertTrue(all(isinstance(value, Decimal) for value in amounts))

        # El legacy no reporta invariantes; el nuevo publica su reconciliación.
        self.assertTrue(self.result.validation.checks)
        self.assertEqual(self.result.validation.warning_codes, ())

    def test_account_header_matches_and_reads_the_holder_alone(self) -> None:
        legacy = dict(_rows(self.legacy_xlsx, "Encabezado")[1:])
        new = dict(_rows(self.new_xlsx, "Encabezado")[1:])

        for legacy_key, new_key in _SHARED_HEADER_FIELDS:
            with self.subTest(field=legacy_key):
                self.assertEqual(str(legacy[legacy_key]), str(new[new_key]))
        # El titular comparte la línea con los rótulos de la cuenta; el nuevo los separa.
        self.assertEqual(new["Titular"], SYNTHETIC_HOLDER)
        self.assertEqual(new["Tipo de cuenta"], "Cuenta de Ahorros CLASICA")
