from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from shutil import rmtree
from unittest import TestCase
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from tests.support.synthetic_bcp_pdf import create_synthetic_bcp_pdf

from statement_worker.domain.errors import (
    ArtifactAlreadyExistsError,
    ArtifactPublicationError,
)
from statement_worker.exporters import (
    BCP_WORKBOOK_SCHEMA_ID,
    BCP_WORKBOOK_SCHEMA_VERSION,
    SpreadsheetCellType,
    SpreadsheetColumn,
    SpreadsheetSheet,
    WorkbookPlan,
    build_bcp_workbook_plan,
    sanitize_spreadsheet_text,
)
from statement_worker.exporters.xlsx_writer import (
    export_workbook_plan,
    verify_workbook_file,
    write_workbook_plan,
)
from statement_worker.extractors.bcp.document_processor import process_bcp_pdf

_FINANCIAL_FORMAT = "#,##0.00;[Red](#,##0.00);-"
_UNSAFE_DESCRIPTION = '=HYPERLINK("http://ficticio")'


def _sample_plan() -> WorkbookPlan:
    return WorkbookPlan(
        schema_id="eecc.statement.test",
        schema_version=1,
        sheets=(
            SpreadsheetSheet(
                name="Movimientos",
                table_name="MovimientosPrueba",
                columns=(
                    SpreadsheetColumn(
                        key="page",
                        header="Página",
                        cell_type=SpreadsheetCellType.INTEGER,
                        width=10,
                        number_format="#,##0",
                    ),
                    SpreadsheetColumn(
                        key="posting_date",
                        header="Fecha proceso",
                        cell_type=SpreadsheetCellType.DATE,
                        width=15,
                        number_format="yyyy-mm-dd",
                    ),
                    SpreadsheetColumn(
                        key="description",
                        header="Descripción",
                        cell_type=SpreadsheetCellType.TEXT,
                        width=48,
                    ),
                    SpreadsheetColumn(
                        key="debit",
                        header="Cargo",
                        cell_type=SpreadsheetCellType.DECIMAL,
                        width=16,
                        number_format=_FINANCIAL_FORMAT,
                    ),
                    SpreadsheetColumn(
                        key="reviewed",
                        header="Revisado",
                        cell_type=SpreadsheetCellType.BOOLEAN,
                        width=12,
                    ),
                ),
                rows=(
                    (
                        1,
                        date(2026, 4, 1),
                        sanitize_spreadsheet_text(_UNSAFE_DESCRIPTION),
                        Decimal("1234.50"),
                        True,
                    ),
                    (2, None, "", None, False),
                ),
            ),
            SpreadsheetSheet(
                name="Validaciones",
                table_name="ValidacionesPrueba",
                columns=(
                    SpreadsheetColumn(
                        key="code",
                        header="Código",
                        cell_type=SpreadsheetCellType.TEXT,
                        width=44,
                    ),
                ),
                rows=(),
            ),
        ),
    )


class XlsxWriterIntegrationTests(TestCase):
    def setUp(self) -> None:
        self.directory = Path.cwd() / f".test_xlsx_writer_{uuid4().hex}"
        self.directory.mkdir()
        self.addCleanup(rmtree, self.directory, True)

    def test_written_file_preserves_types_formats_and_presentation(self) -> None:
        plan = _sample_plan()
        target = self.directory / "statement.xlsx"

        write_workbook_plan(plan, target)

        workbook = load_workbook(filename=target)
        try:
            self.assertEqual(workbook.sheetnames, ["Movimientos", "Validaciones"])
            worksheet = workbook["Movimientos"]
            self.assertEqual(worksheet.freeze_panes, "A2")
            self.assertEqual(worksheet.column_dimensions["C"].width, 48)
            self.assertEqual(
                [cell.value for cell in worksheet[1]],
                ["Página", "Fecha proceso", "Descripción", "Cargo", "Revisado"],
            )
            self.assertEqual(worksheet["A2"].value, 1)
            self.assertEqual(worksheet["B2"].value, datetime(2026, 4, 1))
            self.assertEqual(worksheet["B2"].number_format, "yyyy-mm-dd")
            self.assertEqual(worksheet["C2"].value, "'" + _UNSAFE_DESCRIPTION)
            self.assertEqual(worksheet["C2"].data_type, "s")
            self.assertEqual(worksheet["D2"].value, Decimal("1234.50"))
            self.assertEqual(worksheet["D2"].number_format, _FINANCIAL_FORMAT)
            self.assertIs(worksheet["E2"].value, True)
            self.assertEqual(
                {table.displayName: table.ref for table in worksheet.tables.values()},
                {"MovimientosPrueba": "A1:E3"},
            )

            empty_sheet = workbook["Validaciones"]
            self.assertEqual(empty_sheet.max_row, 1)
            self.assertEqual(empty_sheet.tables, {})
            self.assertEqual(empty_sheet.auto_filter.ref, "A1:A1")
        finally:
            workbook.close()

    def test_written_file_omits_environment_metadata(self) -> None:
        plan = _sample_plan()
        target = self.directory / "metadata.xlsx"

        write_workbook_plan(plan, target)

        workbook = load_workbook(filename=target)
        try:
            self.assertEqual(workbook.properties.creator, "eecc-pdf-worker")
            self.assertEqual(workbook.properties.lastModifiedBy, "eecc-pdf-worker")
            self.assertEqual(workbook.properties.created, datetime(2000, 1, 1))
            self.assertIsNone(workbook.properties.identifier)
            self.assertIsNone(workbook.properties.subject)
        finally:
            workbook.close()
        with ZipFile(target) as archive:
            self.assertNotIn("docProps/custom.xml", archive.namelist())
            self.assertNotIn(str(self.directory), archive.read("docProps/core.xml").decode())

    def test_verification_detects_a_tampered_workbook(self) -> None:
        plan = _sample_plan()
        target = self.directory / "tampered.xlsx"
        write_workbook_plan(plan, target)
        verify_workbook_file(plan, target)

        workbook = load_workbook(filename=target)
        try:
            workbook["Movimientos"]["D2"] = Decimal("1234.51")
            workbook.save(str(target))
        finally:
            workbook.close()

        with self.assertRaises(ArtifactPublicationError):
            verify_workbook_file(plan, target)

    def test_verification_rejects_an_unreadable_file(self) -> None:
        target = self.directory / "broken.xlsx"
        target.write_bytes(b"not a workbook")

        with self.assertRaises(ArtifactPublicationError):
            verify_workbook_file(_sample_plan(), target)

    def test_export_publishes_atomically_and_preserves_the_target(self) -> None:
        plan = _sample_plan()
        target = self.directory / "published.xlsx"

        export_workbook_plan(plan, target)

        self.assertTrue(target.is_file())
        self.assertEqual(list(self.directory.glob(".eecc_artifact_*")), [])
        verify_workbook_file(plan, target)

        published_bytes = target.read_bytes()
        with self.assertRaises(ArtifactAlreadyExistsError):
            export_workbook_plan(plan, target)
        self.assertEqual(target.read_bytes(), published_bytes)

        export_workbook_plan(plan, target, overwrite=True)
        verify_workbook_file(plan, target)

    def test_bcp_statement_reaches_a_reopenable_workbook(self) -> None:
        pdf_path = self.directory / "synthetic-bcp.pdf"
        create_synthetic_bcp_pdf(pdf_path)
        result = process_bcp_pdf(pdf_path, temporary_parent=self.directory)
        plan = build_bcp_workbook_plan(result)
        target = self.directory / "bcp.xlsx"

        export_workbook_plan(plan, target)

        self.assertEqual(plan.schema_id, BCP_WORKBOOK_SCHEMA_ID)
        self.assertEqual(plan.schema_version, BCP_WORKBOOK_SCHEMA_VERSION)
        workbook = load_workbook(filename=target, read_only=True, data_only=True)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["Resumen", "Movimientos", "Control_Paginas", "Validaciones"],
            )
            summary = next(iter(workbook["Resumen"].iter_rows(min_row=2, values_only=True)))
            self.assertEqual(summary[0], result.validation.status.value)
            self.assertEqual(summary[4], len(result.page_metrics))
            self.assertEqual(summary[5], len(result.rows))
            movements = tuple(workbook["Movimientos"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(movements), len(result.rows))
        finally:
            workbook.close()

    def test_verification_detects_type_shape_and_naming_regressions(self) -> None:
        plan = _sample_plan()

        def rename_header(workbook: Workbook) -> None:
            workbook["Movimientos"]["C1"] = "Otro encabezado"

        def append_row(workbook: Workbook) -> None:
            workbook["Movimientos"].append([3, None, "", None, False])

        def text_instead_of_date(workbook: Workbook) -> None:
            workbook["Movimientos"]["B2"] = "2026-04-01"

        def boolean_instead_of_integer(workbook: Workbook) -> None:
            workbook["Movimientos"]["A2"] = True

        def integer_instead_of_boolean(workbook: Workbook) -> None:
            workbook["Movimientos"]["E2"] = 1

        def text_instead_of_amount(workbook: Workbook) -> None:
            workbook["Movimientos"]["D2"] = "1234.50"

        def number_instead_of_text(workbook: Workbook) -> None:
            workbook["Movimientos"]["C2"] = 42

        def drop_sheet(workbook: Workbook) -> None:
            workbook.remove(workbook["Validaciones"])

        tampering = (
            rename_header,
            append_row,
            text_instead_of_date,
            boolean_instead_of_integer,
            integer_instead_of_boolean,
            text_instead_of_amount,
            number_instead_of_text,
            drop_sheet,
        )
        for index, tamper in enumerate(tampering):
            with self.subTest(tamper=tamper.__name__):
                target = self.directory / f"tampered-{index}.xlsx"
                write_workbook_plan(plan, target)
                workbook = load_workbook(filename=target)
                try:
                    tamper(workbook)
                    workbook.save(str(target))
                finally:
                    workbook.close()

                with self.assertRaises(ArtifactPublicationError):
                    verify_workbook_file(plan, target)
