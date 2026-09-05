"""
Convierte los estados de cuenta PDF de la carpeta PDF_movimientos a Excel.

Uso basico:
    py convertir_movimientos_pdf_a_excel.py

Tambien puedes indicar carpetas distintas:
    py convertir_movimientos_pdf_a_excel.py --pdf-dir PDF_movimientos --out-dir Resultados

Dependencias:
    py -m pip install pdfplumber openpyxl

Salida:
    - Un archivo Excel por cada PDF, dentro de Resultados.
    - Un consolidado general: Resultados/movimientos_general.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from contextlib import contextmanager
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Iterable

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ModuleNotFoundError as exc:
    print(f"Falta la dependencia: {exc.name}")
    print("Instala Python 3.10+ y luego ejecuta:")
    print("  py -m pip install pdfplumber openpyxl")
    print("Si usas el comando python en vez de py:")
    print("  python -m pip install pdfplumber openpyxl")
    sys.exit(1)


COLUMNAS_MOVIMIENTOS = [
    "archivo_pdf",
    "periodo",
    "tipo_fila",
    "fecha",
    "fecha_proc",
    "fecha_valor",
    "descripcion",
    "cargo",
    "abono",
    "monto",
    "saldo",
    "moneda",
    "cuenta",
    "pagina",
    "fuente",
    "linea_original",
]

COLUMNAS_ARCHIVOS = [
    "archivo_pdf",
    "periodo",
    "movimientos",
    "excel_individual",
    "estado",
    "mensaje",
]

MONTHS_ES = {
    "ENE": 1,
    "FEB": 2,
    "MAR": 3,
    "ABR": 4,
    "MAY": 5,
    "JUN": 6,
    "JUL": 7,
    "AGO": 8,
    "SET": 9,
    "SEP": 9,
    "OCT": 10,
    "NOV": 11,
    "DIC": 12,
}

DATE_START_RE = re.compile(
    r"^\s*(?P<date>"
    r"\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?"
    r"|"
    r"\d{1,2}\s*[./-]?\s*(?:ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SET|SEP|OCT|NOV|DIC)"
    r"(?:\s*[./-]?\s*\d{2,4})?"
    r")\b",
    re.IGNORECASE,
)

DATE_ANY_RE = re.compile(
    r"(?P<date>"
    r"\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?"
    r"|"
    r"\d{1,2}\s*[./-]?\s*(?:ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SET|SEP|OCT|NOV|DIC)"
    r"(?:\s*[./-]?\s*\d{2,4})?"
    r")\b",
    re.IGNORECASE,
)

AMOUNT_RE = re.compile(
    r"(?<![\w/])"
    r"\(?[-+]?(?:S/\.?\s*)?(?:US\$\s*)?(?:\$?\s*)"
    r"(?:\d{1,3}(?:[,\.\s]\d{3})+|\d+)"
    r"(?:[,.]\d{2})"
    r"\)?"
    r"(?![\w/])",
    re.IGNORECASE,
)

HEADER_PHRASES = (
    "FECHA DESCRIP",
    "FECHA OPER",
    "FECHA DE OPER",
    "DESCRIPCION",
    "CARGOS",
    "ABONOS",
    "SALDO",
    "PAGINA",
)

SKIP_PHRASES = (
    "SALDO ANTERIOR",
    "SALDO INICIAL",
    "SALDO FINAL",
    "TOTAL CARGOS",
    "TOTAL ABONOS",
    "TOTAL MOVIMIENTOS",
    "TOTAL RETIROS",
    "TOTAL DEPOSITOS",
    "RESUMEN",
)

CREDIT_KEYWORDS = (
    "ABONO",
    "DEPOSITO",
    "DEPOSITO EFECTIVO",
    "DEPOSITO EN EFECTIVO",
    "DEPOSITO CHEQUE",
    "DEVOLUCION",
    "EXTORNO",
    "REEMBOLSO",
    "INTERES",
    "TRANSFERENCIA RECIBIDA",
    "TRANSF RECIBIDA",
    "RECIBIDA",
    "RECIBIDO",
    "PAGO DE HABERES",
    "HABERES",
)

DEBIT_KEYWORDS = (
    "COMPRA",
    "RETIRO",
    "COMISION",
    "ITF",
    "CARGO",
    "DEBITO",
    "PAGO",
    "CONSUMO",
    "TRANSFERENCIA ENVIADA",
    "TRANSF ENVIADA",
    "ENVIO",
)

PDF_HEADER = b"%PDF-"
PDF_EOF = b"%%EOF"


@dataclass(frozen=True)
class Periodo:
    texto: str
    anio: int | None
    mes: int | None


@contextmanager
def pdf_legible_temporal(pdf_path: Path):
    """Crea una copia temporal si el PDF viene envuelto con bytes no-PDF."""
    data = pdf_path.read_bytes()
    start = data.find(PDF_HEADER)
    if start == -1:
        raise ValueError("El archivo no contiene una cabecera %PDF valida.")

    eof = data.rfind(PDF_EOF)
    if eof == -1:
        end = len(data)
    else:
        newline = data.find(b"\n", eof)
        end = len(data) if newline == -1 else newline + 1

    if start == 0 and end == len(data):
        yield pdf_path
        return

    cleaned = data[start:end]
    with TemporaryDirectory(prefix="pdf_limpio_") as tmp_dir:
        temp_path = Path(tmp_dir) / pdf_path.name
        temp_path.write_bytes(cleaned)
        yield temp_path


def strip_accents(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in value if not unicodedata.combining(ch))


def normalize_spaces(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def normalized_upper(value: str) -> str:
    return strip_accents(normalize_spaces(value)).upper()


def periodo_desde_nombre(pdf_path: Path) -> Periodo:
    match = re.search(r"EECC(?P<mes>\d{2})(?P<anio>\d{4})", pdf_path.stem, re.IGNORECASE)
    if not match:
        return Periodo("", None, None)

    mes = int(match.group("mes"))
    anio = int(match.group("anio"))
    if 1 <= mes <= 12:
        return Periodo(f"{anio:04d}-{mes:02d}", anio, mes)
    return Periodo("", None, None)


def parse_year(value: str, default_year: int | None) -> int | None:
    if not value:
        return default_year
    year = int(value)
    if year < 100:
        return 2000 + year if year < 70 else 1900 + year
    return year


def parse_fecha(value: str, periodo: Periodo) -> date | None:
    raw = normalized_upper(value)
    raw = raw.replace(".", "/").replace("-", "/")

    numeric = re.match(r"^(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?$", raw)
    if numeric:
        day = int(numeric.group(1))
        month = int(numeric.group(2))
        year = parse_year(numeric.group(3) or "", periodo.anio)
        if year:
            try:
                return date(year, month, day)
            except ValueError:
                return None

    month_names = "|".join(MONTHS_ES)
    text_date = re.match(rf"^(\d{{1,2}})\s*/?\s*({month_names})(?:\s*/?\s*(\d{{2,4}}))?$", raw)
    if text_date:
        day = int(text_date.group(1))
        month = MONTHS_ES[text_date.group(2)]
        year = parse_year(text_date.group(3) or "", periodo.anio)
        if year:
            try:
                return date(year, month, day)
            except ValueError:
                return None

    return None


def parse_importe(value: str) -> Decimal | None:
    raw = normalize_spaces(value)
    if not raw:
        return None

    is_negative = raw.strip().startswith("(") and raw.strip().endswith(")")
    is_negative = is_negative or "-" in raw

    cleaned = raw.upper()
    cleaned = cleaned.replace("S/.", "").replace("S/", "")
    cleaned = cleaned.replace("US$", "").replace("$", "")
    cleaned = cleaned.replace("PEN", "").replace("USD", "")
    cleaned = re.sub(r"[^0-9,.\-]", "", cleaned)
    cleaned = cleaned.replace("-", "")

    if not cleaned:
        return None

    last_dot = cleaned.rfind(".")
    last_comma = cleaned.rfind(",")

    if last_dot > last_comma:
        decimal_sep = "."
        thousands_sep = ","
    elif last_comma > last_dot:
        decimal_sep = ","
        thousands_sep = "."
    else:
        decimal_sep = "." if "." in cleaned else "," if "," in cleaned else ""
        thousands_sep = ""

    if decimal_sep:
        cleaned = cleaned.replace(thousands_sep, "")
        cleaned = cleaned.replace(decimal_sep, ".")

    cleaned = cleaned.replace(" ", "")

    try:
        amount = Decimal(cleaned)
    except InvalidOperation:
        return None

    return -amount if is_negative else amount


def extract_amounts(value: str) -> list[Decimal]:
    amounts: list[Decimal] = []
    for match in AMOUNT_RE.finditer(value):
        parsed = parse_importe(match.group(0))
        if parsed is not None:
            amounts.append(parsed)
    return amounts


def remove_amounts(value: str) -> str:
    return normalize_spaces(AMOUNT_RE.sub(" ", value))


def should_skip_line(value: str, has_date: bool) -> bool:
    norm = normalized_upper(value)
    if not norm:
        return True

    if any(phrase in norm for phrase in SKIP_PHRASES):
        return True

    if not has_date and any(phrase in norm for phrase in HEADER_PHRASES):
        return True

    return False


def looks_credit(description: str) -> bool:
    norm = normalized_upper(description)
    return any(keyword in norm for keyword in CREDIT_KEYWORDS)


def looks_debit(description: str) -> bool:
    norm = normalized_upper(description)
    return any(keyword in norm for keyword in DEBIT_KEYWORDS)


def assign_amounts(amounts: list[Decimal], description: str) -> tuple[Decimal | None, Decimal | None, Decimal | None, Decimal | None]:
    cargo: Decimal | None = None
    abono: Decimal | None = None
    monto: Decimal | None = None
    saldo: Decimal | None = None

    if not amounts:
        return cargo, abono, monto, saldo

    if len(amounts) >= 3:
        cargo = amounts[-3]
        abono = amounts[-2]
        saldo = amounts[-1]
        if cargo == 0:
            cargo = None
        if abono == 0:
            abono = None
        monto = (abono or Decimal("0")) - (cargo or Decimal("0"))
        return cargo, abono, monto, saldo

    movimiento = amounts[0]
    if len(amounts) == 2:
        saldo = amounts[1]

    if movimiento < 0:
        cargo = abs(movimiento)
        monto = movimiento
    elif looks_credit(description):
        abono = movimiento
        monto = movimiento
    elif looks_debit(description):
        cargo = movimiento
        monto = -movimiento
    else:
        # Si el PDF no conserva las columnas cargo/abono, dejamos el importe
        # como monto positivo y la linea original para revision.
        monto = movimiento

    return cargo, abono, monto, saldo


def decimal_to_excel(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value)


def extraer_moneda(texto_pdf: str) -> str:
    norm = normalized_upper(texto_pdf)
    if "US$" in texto_pdf.upper() or "DOLARES" in norm or "USD" in norm:
        return "USD"
    if "S/" in texto_pdf.upper() or "SOLES" in norm or "PEN" in norm:
        return "PEN"
    return ""


def extraer_cuenta(texto_pdf: str) -> str:
    norm = normalize_spaces(texto_pdf)
    patterns = [
        r"(?:CUENTA|CTA\.?)\s*(?:NRO\.?|NO\.?|NUMERO|N\W?)?\s*[:.]?\s*([0-9\- ]{6,})",
        r"(?:NRO\.?|NO\.?|N\W?)\s*(?:DE\s*)?(?:CUENTA|CTA\.?)\s*[:.]?\s*([0-9\- ]{6,})",
        r"(\d{3}-\d{8}-\d-\d{2})",
    ]
    for pattern in patterns:
        match = re.search(pattern, norm, re.IGNORECASE)
        if match:
            return normalize_spaces(match.group(1))
    return ""


def es_fecha_bcp(valor: str) -> bool:
    raw = normalized_upper(valor)
    month_names = "|".join(MONTHS_ES)
    return bool(re.fullmatch(rf"\d{{1,2}}\s*(?:{month_names})", raw))


def obtener_anio_estado_bcp(texto: str, periodo: Periodo) -> int | None:
    norm = normalized_upper(texto)
    match = re.search(
        r"DEL\s+\d{1,2}/\d{1,2}/(\d{2,4})\s+AL\s+\d{1,2}/\d{1,2}/(\d{2,4})",
        norm,
    )
    if match:
        return parse_year(match.group(2), periodo.anio)

    match = re.search(r"\d{1,2}/\d{1,2}/(\d{2,4})", norm)
    if match:
        return parse_year(match.group(1), periodo.anio)

    return periodo.anio


def agrupar_words_por_fila(words: list[dict], tolerancia_y: float = 2.8) -> list[dict]:
    filas: list[dict] = []
    words_ordenados = sorted(words, key=lambda item: (round(item.get("top", 0), 1), item.get("x0", 0)))

    for word in words_ordenados:
        agregado = False
        for fila in filas:
            if abs(fila["top"] - word.get("top", 0)) <= tolerancia_y:
                fila["words"].append(word)
                agregado = True
                break

        if not agregado:
            filas.append({"top": word.get("top", 0), "words": [word]})

    for fila in filas:
        fila["words"] = sorted(fila["words"], key=lambda item: item.get("x0", 0))

    return filas


def clasificar_columna_bcp(x_centro: float) -> str:
    if 30 <= x_centro < 78:
        return "fecha_proc"
    if 78 <= x_centro < 120:
        return "fecha_valor"
    if 120 <= x_centro < 330:
        return "descripcion"
    if 330 <= x_centro < 455:
        return "cargo"
    if 455 <= x_centro < 570:
        return "abono"
    return "fuera_tabla"


def detectar_limites_tabla_bcp(words: list[dict]) -> tuple[float, float]:
    y_inicio = 205.0
    y_fin = 720.0
    tops_header: list[float] = []
    tops_footer: list[float] = []

    for word in words:
        texto = normalized_upper(word.get("text", ""))
        top = float(word.get("top", 0))

        if texto in {"PROC.", "VALOR", "DESCRIPCION", "CARGOS", "DEBE", "ABONOS", "HABER"} and 150 <= top <= 360:
            tops_header.append(top)

        if texto.startswith("ADVERTENCIA") or texto == "MENSAJE":
            if top > 300:
                tops_footer.append(top)

    if tops_header:
        y_inicio = max(tops_header) + 8

    if tops_footer:
        y_fin = min(tops_footer) - 4

    return y_inicio, y_fin


def es_fila_footer_o_ruido_bcp(texto: str) -> bool:
    norm = normalized_upper(texto)
    patrones_ruido = (
        "FECHA PROC",
        "FECHA VALOR",
        "DESCRIPCION",
        "CARGOS / DEBE",
        "ABONOS / HABER",
        "ADVERTENCIA",
        "MENSAJE AL CLIENTE",
        "DEFENSOR DEL CLIENTE",
        "INDECOPI",
        "OFICINAS",
        "311-9898",
        "ESTADO DE CUENTA",
        "CODIGO DE CUENTA",
        "MONEDA",
        "PAGINA",
    )
    return any(pattern in norm for pattern in patrones_ruido)


def monto_neto(cargo: float | None, abono: float | None) -> float | None:
    if cargo is None and abono is None:
        return None
    return float(abono or 0) - float(cargo or 0)


def procesar_fila_bcp(
    fila: dict,
    *,
    pdf_path: Path,
    periodo: Periodo,
    pagina: int,
    anio_estado: int | None,
    moneda: str,
    cuenta: str,
) -> dict | None:
    columnas = {
        "fecha_proc": [],
        "fecha_valor": [],
        "descripcion": [],
        "cargo": [],
        "abono": [],
    }

    for word in fila["words"]:
        texto = normalize_spaces(word.get("text", ""))
        if not texto:
            continue
        x_centro = (float(word.get("x0", 0)) + float(word.get("x1", 0))) / 2
        columna = clasificar_columna_bcp(x_centro)
        if columna in columnas:
            columnas[columna].append(texto)

    texto_completo = normalize_spaces(" ".join(" ".join(values) for values in columnas.values()))
    if not texto_completo or es_fila_footer_o_ruido_bcp(texto_completo):
        return None

    fecha_proc_text = normalize_spaces(" ".join(columnas["fecha_proc"]))
    fecha_valor_text = normalize_spaces(" ".join(columnas["fecha_valor"]))
    descripcion = normalize_spaces(" ".join(columnas["descripcion"]))
    cargo_text = normalize_spaces(" ".join(columnas["cargo"]))
    abono_text = normalize_spaces(" ".join(columnas["abono"]))
    texto_upper = normalized_upper(texto_completo)
    periodo_parse = Periodo(periodo.texto, anio_estado or periodo.anio, periodo.mes)

    row = {
        "archivo_pdf": pdf_path.name,
        "periodo": periodo.texto,
        "tipo_fila": "MOVIMIENTO",
        "fecha": None,
        "fecha_proc": None,
        "fecha_valor": None,
        "descripcion": descripcion,
        "cargo": None,
        "abono": None,
        "monto": None,
        "saldo": None,
        "moneda": moneda,
        "cuenta": cuenta,
        "pagina": pagina,
        "fuente": "bcp_coordenadas",
        "linea_original": texto_completo,
    }

    if "SALDO ANTERIOR" in texto_upper:
        row["tipo_fila"] = "SALDO_ANTERIOR"
        row["descripcion"] = "SALDO ANTERIOR"
        row["saldo"] = decimal_to_excel(parse_importe(abono_text or cargo_text))
        return row

    if "TOTAL MOVIMIENTO" in texto_upper or "TOTAL MOVIMIENTOS" in texto_upper:
        row["tipo_fila"] = "TOTAL_MOVIMIENTO"
        row["descripcion"] = "TOTAL MOVIMIENTO"
        row["cargo"] = decimal_to_excel(parse_importe(cargo_text))
        row["abono"] = decimal_to_excel(parse_importe(abono_text))
        row["monto"] = monto_neto(row["cargo"], row["abono"])
        return row

    if texto_upper == "SALDO" or texto_upper.startswith("SALDO "):
        row["tipo_fila"] = "SALDO"
        row["descripcion"] = "SALDO"
        row["saldo"] = decimal_to_excel(parse_importe(abono_text or cargo_text))
        return row

    if es_fecha_bcp(fecha_proc_text) and es_fecha_bcp(fecha_valor_text):
        fecha_proc = parse_fecha(fecha_proc_text, periodo_parse)
        fecha_valor = parse_fecha(fecha_valor_text, periodo_parse)
        row["fecha_proc"] = fecha_proc
        row["fecha_valor"] = fecha_valor
        row["fecha"] = fecha_valor or fecha_proc
        row["cargo"] = decimal_to_excel(parse_importe(cargo_text))
        row["abono"] = decimal_to_excel(parse_importe(abono_text))
        row["monto"] = monto_neto(row["cargo"], row["abono"])
        return row

    if descripcion:
        row["tipo_fila"] = "CONTINUACION"
        row["cargo"] = decimal_to_excel(parse_importe(cargo_text))
        row["abono"] = decimal_to_excel(parse_importe(abono_text))
        row["monto"] = monto_neto(row["cargo"], row["abono"])
        return row

    return None


def unir_continuaciones_bcp(rows: list[dict]) -> list[dict]:
    final_rows: list[dict] = []

    for row in rows:
        if row.get("tipo_fila") != "CONTINUACION":
            final_rows.append(row)
            continue

        if not final_rows:
            continue

        descripcion_extra = normalize_spaces(row.get("descripcion", ""))
        if descripcion_extra:
            final_rows[-1]["descripcion"] = normalize_spaces(
                f"{final_rows[-1].get('descripcion', '')} | {descripcion_extra}"
            )

        for column in ("cargo", "abono", "monto", "saldo"):
            if final_rows[-1].get(column) is None and row.get(column) is not None:
                final_rows[-1][column] = row.get(column)

    return final_rows


def extraer_movimientos_bcp_por_coordenadas(pdf_path: Path, password: str | None = None) -> list[dict]:
    periodo = periodo_desde_nombre(pdf_path)
    rows: list[dict] = []
    moneda = ""
    cuenta = ""
    anio_estado = periodo.anio

    with pdf_legible_temporal(pdf_path) as readable_pdf_path:
        with pdfplumber.open(readable_pdf_path, password=password or "") as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                if page_number == 1:
                    first_page_text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
                    moneda = extraer_moneda(first_page_text)
                    cuenta = extraer_cuenta(first_page_text)
                    anio_estado = obtener_anio_estado_bcp(first_page_text, periodo)

                page_tabla = page.crop((25, 180, min(575, page.width), min(740, page.height)))
                words = page_tabla.extract_words(x_tolerance=1, y_tolerance=2, keep_blank_chars=False)
                y_inicio, y_fin = detectar_limites_tabla_bcp(words)
                words_tabla = [word for word in words if y_inicio <= float(word.get("top", 0)) <= y_fin]
                filas = agrupar_words_por_fila(words_tabla, tolerancia_y=2.8)

                for fila in filas:
                    row = procesar_fila_bcp(
                        fila,
                        pdf_path=pdf_path,
                        periodo=periodo,
                        pagina=page_number,
                        anio_estado=anio_estado,
                        moneda=moneda,
                        cuenta=cuenta,
                    )
                    if row:
                        rows.append(row)

    return unir_continuaciones_bcp(rows)


def movimiento_desde_linea(
    line: str,
    *,
    pdf_path: Path,
    periodo: Periodo,
    pagina: int,
    fuente: str,
    moneda: str,
    cuenta: str,
) -> dict | None:
    raw = normalize_spaces(line)
    match = DATE_START_RE.match(raw)
    if not match:
        return None

    fecha = parse_fecha(match.group("date"), periodo)
    if fecha is None:
        return None

    if should_skip_line(raw, has_date=True):
        return None

    body = raw[match.end() :].strip()
    amounts = extract_amounts(body)
    if not amounts:
        return None

    descripcion = remove_amounts(body)
    descripcion = DATE_ANY_RE.sub(" ", descripcion, count=1)
    descripcion = normalize_spaces(descripcion)

    if not descripcion:
        return None

    cargo, abono, monto, saldo = assign_amounts(amounts, descripcion)

    return {
        "archivo_pdf": pdf_path.name,
        "periodo": periodo.texto,
        "tipo_fila": "MOVIMIENTO",
        "fecha": fecha,
        "fecha_proc": fecha,
        "fecha_valor": fecha,
        "descripcion": descripcion,
        "cargo": decimal_to_excel(cargo),
        "abono": decimal_to_excel(abono),
        "monto": decimal_to_excel(monto),
        "saldo": decimal_to_excel(saldo),
        "moneda": moneda,
        "cuenta": cuenta,
        "pagina": pagina,
        "fuente": fuente,
        "linea_original": raw,
    }


def movimiento_desde_fila_tabla(
    cells: Iterable[str | None],
    *,
    pdf_path: Path,
    periodo: Periodo,
    pagina: int,
    moneda: str,
    cuenta: str,
) -> dict | None:
    cleaned_cells = [normalize_spaces(cell) for cell in cells if normalize_spaces(cell)]
    if not cleaned_cells:
        return None

    raw = normalize_spaces(" ".join(cleaned_cells))
    date_match: re.Match[str] | None = None
    date_cell_index: int | None = None

    for index, cell in enumerate(cleaned_cells[:4]):
        date_match = DATE_START_RE.match(cell) or DATE_ANY_RE.search(cell)
        if date_match:
            date_cell_index = index
            break

    if not date_match:
        return movimiento_desde_linea(
            raw,
            pdf_path=pdf_path,
            periodo=periodo,
            pagina=pagina,
            fuente="texto_tabla",
            moneda=moneda,
            cuenta=cuenta,
        )

    fecha = parse_fecha(date_match.group("date"), periodo)
    if fecha is None or should_skip_line(raw, has_date=True):
        return None

    amounts = extract_amounts(raw)
    if not amounts:
        return None

    description_parts: list[str] = []
    for index, cell in enumerate(cleaned_cells):
        text = remove_amounts(cell)
        if index == date_cell_index and date_match:
            text = text.replace(date_match.group("date"), " ")
        text = DATE_ANY_RE.sub(" ", text)
        text = normalize_spaces(text)
        if text and not should_skip_line(text, has_date=False):
            description_parts.append(text)

    descripcion = normalize_spaces(" ".join(description_parts))
    if not descripcion:
        return None

    cargo, abono, monto, saldo = assign_amounts(amounts, descripcion)

    return {
        "archivo_pdf": pdf_path.name,
        "periodo": periodo.texto,
        "tipo_fila": "MOVIMIENTO",
        "fecha": fecha,
        "fecha_proc": fecha,
        "fecha_valor": fecha,
        "descripcion": descripcion,
        "cargo": decimal_to_excel(cargo),
        "abono": decimal_to_excel(abono),
        "monto": decimal_to_excel(monto),
        "saldo": decimal_to_excel(saldo),
        "moneda": moneda,
        "cuenta": cuenta,
        "pagina": pagina,
        "fuente": "tabla_pdf",
        "linea_original": raw,
    }


def dedupe_movimientos(rows: list[dict]) -> list[dict]:
    seen: set[tuple] = set()
    unique_rows: list[dict] = []

    for row in rows:
        key = (
            row.get("archivo_pdf"),
            row.get("tipo_fila"),
            row.get("fecha"),
            row.get("fecha_proc"),
            row.get("fecha_valor"),
            normalized_upper(str(row.get("descripcion", ""))),
            row.get("cargo"),
            row.get("abono"),
            row.get("monto"),
            row.get("saldo"),
            row.get("pagina"),
        )
        if key in seen:
            continue
        seen.add(key)
        unique_rows.append(row)

    return unique_rows


def extraer_movimientos_pdf(pdf_path: Path, password: str | None = None) -> list[dict]:
    bcp_rows = extraer_movimientos_bcp_por_coordenadas(pdf_path, password=password)
    if any(row.get("tipo_fila") == "MOVIMIENTO" for row in bcp_rows):
        return dedupe_movimientos(bcp_rows)

    periodo = periodo_desde_nombre(pdf_path)
    rows: list[dict] = []
    text_chunks: list[str] = []

    table_settings_options = [
        {
            "vertical_strategy": "lines",
            "horizontal_strategy": "lines",
            "snap_tolerance": 3,
            "join_tolerance": 3,
            "intersection_tolerance": 5,
        },
        {
            "vertical_strategy": "text",
            "horizontal_strategy": "text",
            "snap_tolerance": 3,
            "join_tolerance": 3,
            "intersection_tolerance": 5,
            "text_x_tolerance": 2,
            "text_y_tolerance": 3,
        },
    ]

    with pdf_legible_temporal(pdf_path) as readable_pdf_path:
        with pdfplumber.open(readable_pdf_path, password=password or "") as pdf:
            for page in pdf.pages:
                page_text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
                text_chunks.append(page_text)

        full_text = "\n".join(text_chunks)
        moneda = extraer_moneda(full_text)
        cuenta = extraer_cuenta(full_text)

        with pdfplumber.open(readable_pdf_path, password=password or "") as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                for table_settings in table_settings_options:
                    try:
                        tables = page.extract_tables(table_settings=table_settings) or []
                    except Exception:
                        tables = []

                    for table in tables:
                        for table_row in table:
                            movement = movimiento_desde_fila_tabla(
                                table_row,
                                pdf_path=pdf_path,
                                periodo=periodo,
                                pagina=page_number,
                                moneda=moneda,
                                cuenta=cuenta,
                            )
                            if movement:
                                rows.append(movement)

                page_text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
                for line in page_text.splitlines():
                    movement = movimiento_desde_linea(
                        line,
                        pdf_path=pdf_path,
                        periodo=periodo,
                        pagina=page_number,
                        fuente="texto_pdf",
                        moneda=moneda,
                        cuenta=cuenta,
                    )
                    if movement:
                        rows.append(movement)

    return dedupe_movimientos(rows)


def setup_sheet(ws, columns: list[str], rows: list[dict], table_name: str) -> None:
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, None) for column in columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")

    for col_index, column_name in enumerate(columns, start=1):
        letter = get_column_letter(col_index)
        max_len = len(column_name)
        for cell in ws[letter]:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)

        if column_name in {"fecha", "fecha_proc", "fecha_valor"}:
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        elif column_name in {"cargo", "abono", "monto", "saldo"}:
            ws.column_dimensions[letter].width = 14
            for cell in ws[letter][1:]:
                cell.number_format = '#,##0.00;[Red]-#,##0.00;0.00'
        elif column_name == "linea_original":
            ws.column_dimensions[letter].width = 80
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if ws.max_row >= 2 and ws.max_column >= 1:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=table_name, ref=ref)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)


def guardar_excel_individual(rows: list[dict], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    setup_sheet(ws, COLUMNAS_MOVIMIENTOS, rows, "TablaMovimientos")
    wb.save(output_path)


def resumen_mensual(rows: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = defaultdict(
        lambda: {
            "periodo": "",
            "movimientos": 0,
            "total_cargos": 0.0,
            "total_abonos": 0.0,
            "movimiento_neto": 0.0,
            "saldo_final": None,
        }
    )

    for row in rows:
        periodo = row.get("periodo") or "sin_periodo"
        item = grouped[periodo]
        item["periodo"] = periodo

        if row.get("tipo_fila") == "MOVIMIENTO":
            item["movimientos"] += 1
            item["total_cargos"] += float(row.get("cargo") or 0)
            item["total_abonos"] += float(row.get("abono") or 0)
            item["movimiento_neto"] += float(row.get("monto") or 0)

        if row.get("saldo") is not None:
            item["saldo_final"] = row.get("saldo")

    return [grouped[key] for key in sorted(grouped)]


def guardar_excel_general(rows: list[dict], file_statuses: list[dict], output_path: Path) -> None:
    wb = Workbook()

    ws_mov = wb.active
    ws_mov.title = "Movimientos"
    setup_sheet(ws_mov, COLUMNAS_MOVIMIENTOS, rows, "TablaGeneral")

    ws_resumen = wb.create_sheet("Resumen_mensual")
    resumen_rows = resumen_mensual(rows)
    setup_sheet(
        ws_resumen,
        ["periodo", "movimientos", "total_cargos", "total_abonos", "movimiento_neto", "saldo_final"],
        resumen_rows,
        "TablaResumenMensual",
    )

    ws_archivos = wb.create_sheet("Archivos")
    setup_sheet(ws_archivos, COLUMNAS_ARCHIVOS, file_statuses, "TablaArchivos")

    wb.save(output_path)


def listar_pdfs(pdf_dir: Path) -> list[Path]:
    pdfs = [path for path in pdf_dir.iterdir() if path.is_file() and path.suffix.lower() == ".pdf"]
    return sorted(pdfs, key=lambda path: path.name.lower())


def convertir_todos(pdf_dir: Path, out_dir: Path, password: str | None = None) -> tuple[list[dict], list[dict]]:
    if not pdf_dir.exists():
        raise FileNotFoundError(f"No existe la carpeta de PDF: {pdf_dir}")

    out_dir.mkdir(parents=True, exist_ok=True)
    pdfs = listar_pdfs(pdf_dir)
    if not pdfs:
        raise FileNotFoundError(f"No encontre archivos PDF en: {pdf_dir}")

    all_rows: list[dict] = []
    file_statuses: list[dict] = []

    for pdf_path in pdfs:
        periodo = periodo_desde_nombre(pdf_path)
        individual_path = out_dir / f"{pdf_path.stem}.xlsx"

        print(f"Procesando: {pdf_path.name}")
        try:
            rows = extraer_movimientos_pdf(pdf_path, password=password)
            guardar_excel_individual(rows, individual_path)
            all_rows.extend(rows)
            movement_count = sum(1 for row in rows if row.get("tipo_fila") == "MOVIMIENTO")
            file_statuses.append(
                {
                    "archivo_pdf": pdf_path.name,
                    "periodo": periodo.texto,
                    "movimientos": movement_count,
                    "excel_individual": str(individual_path),
                    "estado": "OK",
                    "mensaje": "",
                }
            )
            print(f"  -> {movement_count} movimientos: {individual_path.name}")
        except Exception as exc:
            file_statuses.append(
                {
                    "archivo_pdf": pdf_path.name,
                    "periodo": periodo.texto,
                    "movimientos": 0,
                    "excel_individual": "",
                    "estado": "ERROR",
                    "mensaje": str(exc),
                }
            )
            print(f"  -> ERROR: {exc}")

    general_path = out_dir / "movimientos_general.xlsx"
    guardar_excel_general(all_rows, file_statuses, general_path)
    print(f"\nConsolidado creado: {general_path}")
    return all_rows, file_statuses


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convierte PDFs de estados de cuenta a Excel individual y consolidado."
    )
    parser.add_argument(
        "--pdf-dir",
        default="PDF_MOVIMIENTOS",
        help="Carpeta donde estan los PDF. Por defecto: PDF_MOVIMIENTOS",
    )
    parser.add_argument(
        "--out-dir",
        default="RESULTADOS",
        help="Carpeta donde se guardan los Excel. Por defecto: RESULTADOS",
    )
    parser.add_argument(
        "--password",
        default=None,
        help="Password de los PDF, si todos usan la misma clave.",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    base_dir = Path(__file__).resolve().parent
    pdf_dir = Path(args.pdf_dir)
    out_dir = Path(args.out_dir)

    if not pdf_dir.is_absolute():
        pdf_dir = base_dir / pdf_dir
    if not out_dir.is_absolute():
        out_dir = base_dir / out_dir

    try:
        rows, statuses = convertir_todos(pdf_dir, out_dir, password=args.password)
    except Exception as exc:
        print(f"Error: {exc}")
        return 1

    ok_files = sum(1 for item in statuses if item["estado"] == "OK")
    error_files = sum(1 for item in statuses if item["estado"] == "ERROR")
    movement_count = sum(1 for row in rows if row.get("tipo_fila") == "MOVIMIENTO")
    print("\nResumen:")
    print(f"  PDFs procesados correctamente: {ok_files}")
    print(f"  PDFs con error: {error_files}")
    print(f"  Movimientos consolidados: {movement_count}")
    return 0 if error_files == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
