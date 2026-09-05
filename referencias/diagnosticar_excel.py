"""
Diagnosticar qué está mal con el archivo Excel
"""
import sys
from pathlib import Path

archivo = Path("2026_JUNIO_BCPxd.xlsx")

print(f"📁 Archivo: {archivo.name}")
print(f"📊 Existe: {archivo.exists()}")
print(f"📏 Tamaño: {archivo.stat().st_size if archivo.exists() else 'N/A'} bytes")

if archivo.exists():
    # Verificar si es un archivo ZIP válido (los .xlsx son ZIPs)
    try:
        import zipfile
        if zipfile.is_zipfile(archivo):
            print("✅ Es un archivo ZIP válido")
            with zipfile.ZipFile(archivo, 'r') as z:
                print(f"📦 Contiene {len(z.namelist())} archivos internos")
                # Ver si tiene la estructura correcta de Excel
                if '_rels/.rels' in z.namelist():
                    print("✅ Estructura Excel válida encontrada")
                else:
                    print("⚠️ Estructura Excel incompleta")
        else:
            print("❌ NO es un archivo ZIP válido (archivo Excel corrupto)")
    except Exception as e:
        print(f"❌ Error verificando: {e}")
    
    # Intentar leer con openpyxl si está disponible
    try:
        from openpyxl import load_workbook
        wb = load_workbook(archivo)
        print(f"✅ Abierto correctamente con openpyxl")
        print(f"📋 Hojas: {wb.sheetnames}")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"   └─ '{sheet_name}': {ws.max_row} filas × {ws.max_column} columnas")
    except ModuleNotFoundError:
        print("⚠️ openpyxl no instalado")
    except Exception as e:
        print(f"❌ Error abriendo con openpyxl: {e}")
